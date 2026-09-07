"""Build matrix XLSX from cleaned Almazar Document Intelligence JSON."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from project_paths import OUTPUT_DIR, PROCESSING_DIR

CURRENCY = "EUR"

COST_NAME_ROW = 1
APPLY_IF_ROW = 2
RATE_BY_ROW = 3
BRACKET_ROW = 4
COLUMN_HEADER_ROW = 5
DATA_START_ROW = 6

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
BOLD = Font(bold=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SHIPMENT_HEADERS = [
    "Origin Postal Code Zone",
    "Destination Postal Code Zone",
]

ORIGIN_BARCELONA = "Barcelona"
ORIGIN_TENERIFE = "Tenerife"
ORIGIN_MAINLAND_SPAIN = "Mainland Spain"
INSURANCE_FCL_40_LABEL = (
    "l[100% - Shipment Insurance (Carga EN CELRA/CORAL'40)%]"
)
EMERGENCY_BUNKER_NOTE = "TBD"

TENERIFE_DIST_BRACKETS = (
    "MIN",
    "<=150",
    "<=3000",
    ">3000",
)

FCL_SIZE_BRACKETS = ("20 pies", "40 pies")

ISLAND_DESTINATIONS = (
    "Las Palmas",
    "Lanzarote",
    "Fuerteventura",
    "La Palma",
    "La Gomera",
    "El Hierro",
)

GRUPAJE_DESTINATIONS = ("Las Palmas", "Lanzarote", "Tenerife")

FCL_ORIGINS = ("Las Palmas", "Tenerife")


@dataclass(frozen=True)
class CostColumnSpec:
    bracket_label: str
    rate_unit: str = "Flat"


@dataclass
class CostBlock:
    title: str
    apply_if: str
    rate_by: str
    columns: list[CostColumnSpec] = field(default_factory=list)
    uses_shared_currency: bool = False


@dataclass
class MatrixRow:
    shipment: dict[str, Any]
    costs: dict[tuple[str, str], float | str | None] = field(default_factory=dict)
    text_costs: dict[tuple[str, str], str] = field(default_factory=dict)


@dataclass
class SkippedItem:
    source: str
    label: str
    detail: str


def _norm_key(text: object) -> str:
    raw = unicodedata.normalize("NFKD", str(text or ""))
    raw = "".join(ch for ch in raw if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", raw.casefold())


def _row_value(row: dict[str, Any], key: str) -> Any:
    if key in row:
        return row[key]
    colon_key = f"{key}:"
    if colon_key in row:
        return row[colon_key]
    return None


def _parse_euro(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("€", "").replace(" ", "").strip()
    if text.endswith("%"):
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        whole, frac = text.split(",", 1)
        if len(frac) == 2:
            text = f"{whole}.{frac}"
        elif len(frac) == 3:
            if whole in ("0", "") or (whole.isdigit() and int(whole) < 10):
                text = f"{whole}.{frac}"
            else:
                text = f"{whole}.{frac[:2]}"
        else:
            text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _place_label(place: str) -> str:
    return place.strip()


def _load_fields(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        payload = json.load(handle)
    if not isinstance(payload, dict):
        raise ValueError("Expected JSON root object.")
    if "analyzeResult" in payload:
        from clean_di_json import clean_analyze_payload

        return clean_analyze_payload(payload)["documents"][0]["fields"]
    documents = payload.get("documents")
    if isinstance(documents, list) and documents:
        fields = documents[0].get("fields")
        if isinstance(fields, dict):
            return fields
    raise ValueError("JSON does not contain document fields.")


def _iter_sections(
    rows: list[dict[str, Any]], *, name_key: str = "RateName"
) -> Iterator[tuple[str, list[dict[str, Any]]]]:
    current_name: str | None = None
    buffer: list[dict[str, Any]] = []

    def flush() -> Iterator[tuple[str, list[dict[str, Any]]]]:
        nonlocal buffer, current_name
        if current_name is not None and buffer:
            yield current_name, buffer
        buffer = []

    for row in rows:
        rate_name = row.get(name_key)
        if rate_name:
            yield from flush()
            current_name = str(rate_name).strip()
            buffer = []
            continue
        if current_name is not None:
            buffer.append(row)
    yield from flush()


def _section_rows(
    rows: list[dict[str, Any]], *name_fragments: str
) -> list[dict[str, Any]]:
    targets = {_norm_key(fragment) for fragment in name_fragments}
    for section_name, section_rows in _iter_sections(rows):
        key = _norm_key(section_name)
        if any(target in key or key in target for target in targets):
            return section_rows
    return []


def _match_origin(origin: object, expected: str) -> bool:
    return _norm_key(origin) == _norm_key(expected)


def _match_weight(weight: object, expected: str) -> bool:
    return _norm_key(weight) == _norm_key(expected)


def _parse_fcl_table(section_rows: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, float | None]]:
    """Map (origin, size) -> {cost1: Arrastres/Flete, cost2: BAFF, cost3: Tarifa 2026}."""
    table: dict[tuple[str, str], dict[str, float | None]] = {}
    for row in section_rows:
        origin = _row_value(row, "Origin")
        size = _row_value(row, "Value1")
        if not origin or not size:
            continue
        if _norm_key(origin) == _norm_key("Origen"):
            continue
        size_norm = _norm_key(size)
        if "20" in size_norm:
            size_key = "20 pies"
        elif "40" in size_norm:
            size_key = "40 pies"
        else:
            continue
        origin_key = str(origin).strip()
        table[(origin_key, size_key)] = {
            "cost1": _parse_euro(_row_value(row, "Cost1")),
            "cost2": _parse_euro(_row_value(row, "Cost2")),
            "cost3": _parse_euro(_row_value(row, "Cost3")),
        }
    return table


def _parse_grupajes(section_rows: list[dict[str, Any]]) -> dict[str, dict[str, float | None]]:
    rates: dict[str, dict[str, float | None]] = {}
    for row in section_rows:
        origin = _row_value(row, "Origin")
        if not origin:
            continue
        name = str(origin).strip()
        if _norm_key(name) in {
            _norm_key("Minimo"),
            _norm_key("Entrega /Reexp en destino"),
            _norm_key("Recargo Transporte 22/06/22"),
            _norm_key("Fuel (Fes)"),
            _norm_key("Despacho Origen"),
            _norm_key("Despacho Destino"),
        } or "seguro" in _norm_key(name):
            continue
        per_kg = _parse_euro(_row_value(row, "Cost2"))
        minimum = _parse_euro(_row_value(row, "Cost3"))
        if per_kg is None and minimum is None:
            continue
        rates[name] = {"per_kg": per_kg, "minimum": minimum}
    return rates


def _parse_tenerife_distribution(
    section_rows: list[dict[str, Any]],
) -> dict[str, float | None]:
    mapping: dict[str, float | None] = {}
    for row in section_rows:
        weight = _row_value(row, "Weight")
        price = _parse_euro(_row_value(row, "Price"))
        if not weight or price is None:
            continue
        w = str(weight).strip()
        key = _norm_key(w)
        if key in {_norm_key("MIN"), _norm_key("Minimo")} or (
            "minimo" in key and "150" not in key and "3000" not in key
        ):
            mapping["MIN"] = price
        elif w.startswith(">") or ("mas" in key and "3000" in key):
            mapping[">3000"] = price
        elif w.startswith("<=") and "150" in key:
            mapping["<=150"] = price
        elif w.startswith("<=") and "3000" in key:
            mapping["<=3000"] = price
        elif "150" in key and "3000" not in key and "mas" not in key:
            mapping["<=150"] = price
        elif "3000" in key and "mas" not in key:
            mapping["<=3000"] = price
    return mapping


def _parse_wh_distribution(
    section_rows: list[dict[str, Any]],
) -> dict[str, dict[str, float | None]]:
    rates: dict[str, dict[str, float | None]] = {}
    for row in section_rows:
        dest = _row_value(row, "Weight")
        per_kg = _parse_euro(_row_value(row, "Price"))
        minimum = _parse_euro(_row_value(row, "MinPrice"))
        if not dest:
            continue
        if per_kg is None and minimum is None:
            continue
        rates[str(dest).strip()] = {"per_kg": per_kg, "minimum": minimum}
    return rates


def _excel_number_format(value: float) -> str:
    """Preserve parsed decimals in Excel without rounding to two places."""
    text = f"{value:.10f}".rstrip("0").rstrip(".")
    decimals = len(text.split(".")[1]) if "." in text else 0
    if decimals <= 2:
        return "0.00" if decimals == 2 else "0.0" if decimals == 1 else "0"
    return "0." + ("0" * decimals)


def _parse_otros_conceptos(section_rows: list[dict[str, Any]]) -> dict[str, str | float | None]:
    data: dict[str, str | float | None] = {}
    for row in section_rows:
        weight = _row_value(row, "Weight")
        price = _row_value(row, "Price")
        if not weight:
            continue
        key = _norm_key(weight)
        if "despacho" in key:
            data["customs_combined"] = _parse_euro(price) or str(price).strip()
        elif "seguro" in key:
            data["insurance"] = str(price).strip() if price else None
    return data


def _size_title_label(size: str) -> str:
    return "20'" if "20" in _norm_key(size) else "40'"


def _fcl_service_blocks(
    service_label: str,
    sizes: tuple[str, ...],
    *,
    default_apply: str,
) -> list[CostBlock]:
    blocks: list[CostBlock] = []
    for size in sizes:
        size_label = _size_title_label(size)
        blocks.append(
            CostBlock(
                title=f"Transport cost ({service_label}, {size_label})",
                apply_if=default_apply,
                rate_by=(
                    f"Rate by: Arrastres/Flete ({service_label}, {size_label})"
                ),
                columns=[CostColumnSpec("Flat", "Flat")],
            )
        )
        blocks.append(
            CostBlock(
                title=f"BAF ({service_label}, {size_label})",
                apply_if=default_apply,
                rate_by=f"Rate by: BAFF ({service_label}, {size_label})",
                columns=[CostColumnSpec("Flat", "Flat")],
            )
        )
    return blocks


def _blocks_by_title(cost_blocks: list[CostBlock]) -> dict[str, CostBlock]:
    return {block.title: block for block in cost_blocks}


def _set_flat_cost(
    costs: dict[tuple[str, str], float | str | None],
    block: CostBlock | None,
    value: float | str | None,
) -> None:
    if block is None or not _has_cost(value):
        return
    costs[cost_key(block, block.columns[0])] = value


def _assign_fcl_service_costs(
    costs: dict[tuple[str, str], float | str | None],
    blocks_by_title: dict[str, CostBlock],
    table: dict[tuple[str, str], dict[str, float | None]],
    service_label: str,
    origin_name: str,
    sizes: tuple[str, ...],
) -> None:
    for size in sizes:
        row = table.get((origin_name, size))
        if not row:
            continue
        size_label = _size_title_label(size)
        transport_block = blocks_by_title.get(
            f"Transport cost ({service_label}, {size_label})"
        )
        baf_block = blocks_by_title.get(f"BAF ({service_label}, {size_label})")
        _set_flat_cost(costs, transport_block, row.get("cost1"))
        _set_flat_cost(costs, baf_block, row.get("cost2"))


def _cost_blocks() -> list[CostBlock]:
    default_apply = "Apply if: Applies in all items"
    return [
        CostBlock(
            title="Transport cost (Distribución en la Isla de Tenerife)",
            apply_if=default_apply,
            rate_by="Rate by: Weight (Distribución Isla Tenerife)",
            columns=[
                CostColumnSpec(label, "Flat") for label in TENERIFE_DIST_BRACKETS
            ],
            uses_shared_currency=True,
        ),
        CostBlock(
            title="Transport cost (Envios)",
            apply_if=default_apply,
            rate_by="Rate by: Per kg (Distribución desde almacén Tenerife)",
            columns=[
                CostColumnSpec("MIN", "Flat"),
                CostColumnSpec("Flat", "Per kg"),
            ],
            uses_shared_currency=True,
        ),
        CostBlock(
            title=(
                "Origin Customs Clearance (ALL) + Destination Customs Clearance "
                "(Grupajes (Despacho Destino))"
            ),
            apply_if=default_apply,
            rate_by="Rate by: Per shipment (Otros conceptos)",
            columns=[CostColumnSpec("Flat", "Flat")],
        ),
        CostBlock(
            title="Shipment Insurance (Cargas Completas (Seguro))",
            apply_if=default_apply,
            rate_by="Rate by: Percent of freight",
            columns=[CostColumnSpec("Flat", "Flat")],
        ),
        *_fcl_service_blocks("Cargas EN CORAL", FCL_SIZE_BRACKETS, default_apply=default_apply),
        *_fcl_service_blocks("Cargas EN CELRA", FCL_SIZE_BRACKETS, default_apply=default_apply),
        *_fcl_service_blocks(
            "CARGA EN CELRA / CORAL",
            ("40 pies",),
            default_apply=default_apply,
        ),
        CostBlock(
            title="Transport cost (Grupajes)",
            apply_if=default_apply,
            rate_by="Rate by: Per kg (Zona Destino / Grupajes, Tarifa 2025)",
            columns=[
                CostColumnSpec("Per kg", "Per kg"),
                CostColumnSpec("Minimum", "Flat"),
            ],
            uses_shared_currency=True,
        ),
        CostBlock(
            title="Emergency Bunker Surcharge",
            apply_if=default_apply,
            rate_by="Rate by: TBD",
            columns=[CostColumnSpec("Flat", "Flat")],
        ),
        CostBlock(
            title="Shipment Insurance (Carga EN CELRA/CORAL'40)",
            apply_if=default_apply,
            rate_by="Rate by: Percent of freight",
            columns=[CostColumnSpec("Flat", "Flat")],
        ),
    ]


def cost_key(block: CostBlock, spec: CostColumnSpec) -> tuple[str, str]:
    return (block.title, spec.bracket_label)


def block_column_width(block: CostBlock) -> int:
    if block.uses_shared_currency:
        return 1 + len(block.columns)
    return 2 * len(block.columns)


def _has_cost(value: float | str | None) -> bool:
    return value is not None and value != ""


def _block_has_row_data(matrix_row: MatrixRow, block: CostBlock) -> bool:
    for spec in block.columns:
        key = cost_key(block, spec)
        if key in matrix_row.text_costs:
            return True
        if _has_cost(matrix_row.costs.get(key)):
            return True
    return False


def _build_lane_rows() -> list[tuple[str, str]]:
    lanes: list[tuple[str, str]] = []
    for dest in GRUPAJE_DESTINATIONS:
        lanes.append((ORIGIN_BARCELONA, _place_label(dest)))
    for dest in ISLAND_DESTINATIONS:
        lanes.append((ORIGIN_TENERIFE, _place_label(dest)))
    lanes.append((ORIGIN_TENERIFE, ORIGIN_TENERIFE))
    for port in FCL_ORIGINS:
        lanes.append((ORIGIN_MAINLAND_SPAIN, _place_label(port)))
    return lanes


def build_matrix(
    fields: dict[str, Any],
) -> tuple[list[MatrixRow], list[CostBlock], list[SkippedItem]]:
    main = fields.get("MainCosts") or []
    main2 = fields.get("MainCosts2") or []
    skipped: list[SkippedItem] = []

    coral_rows = _section_rows(main, "Carga en Coral")
    celra_rows = _section_rows(main, "Carga en Celra")
    combo_rows = _section_rows(main, "Carga en Celra + Coral", "Celra + Coral")
    grupaje_rows = _section_rows(main, "Zona Destino/Grupajes", "Grupajes")

    tenerife_dist_rows = _section_rows(
        main2, "Disstribucion isla tenerife", "Distribucion isla tenerife"
    )
    wh_dist_rows = _section_rows(
        main2,
        "Distribution desde almacen de tenerife",
        "Distribución desde almacén de Tenerife",
    )
    otros_rows = _section_rows(main2, "Otros conceptos")

    coral = _parse_fcl_table(coral_rows)
    celra = _parse_fcl_table(celra_rows)
    combo = _parse_fcl_table(combo_rows)
    grupajes = _parse_grupajes(grupaje_rows)
    tenerife_dist = _parse_tenerife_distribution(tenerife_dist_rows)
    wh_dist = _parse_wh_distribution(wh_dist_rows)
    otros = _parse_otros_conceptos(otros_rows)

    customs = otros.get("customs_combined")
    insurance_cc = otros.get("insurance")

    cost_blocks = _cost_blocks()
    blocks_by_title = _blocks_by_title(cost_blocks)
    matrix_rows: list[MatrixRow] = []

    customs_block = blocks_by_title[
        "Origin Customs Clearance (ALL) + Destination Customs Clearance "
        "(Grupajes (Despacho Destino))"
    ]
    insurance_cc_block = blocks_by_title["Shipment Insurance (Cargas Completas (Seguro))"]
    grupajes_block = blocks_by_title["Transport cost (Grupajes)"]
    insurance_fcl_block = blocks_by_title["Shipment Insurance (Carga EN CELRA/CORAL'40)"]

    for origin_zone, dest_zone in _build_lane_rows():
        shipment = {
            "Origin Postal Code Zone": origin_zone,
            "Destination Postal Code Zone": dest_zone,
        }
        costs: dict[tuple[str, str], float | str | None] = {}
        text_costs: dict[tuple[str, str], str] = {}

        dest_name = dest_zone
        port_name = dest_zone

        # Tenerife island distribution (Distribución en la Isla de Tenerife)
        if origin_zone == ORIGIN_TENERIFE and dest_zone == ORIGIN_TENERIFE:
            block = cost_blocks[0]
            for bracket in TENERIFE_DIST_BRACKETS:
                costs[cost_key(block, CostColumnSpec(bracket))] = tenerife_dist.get(
                    bracket
                )

        # Warehouse distribution / envíos (Distribución desde almacén de Tenerife)
        if origin_zone == ORIGIN_TENERIFE and dest_zone in ISLAND_DESTINATIONS:
            block = cost_blocks[1]
            for island in ISLAND_DESTINATIONS:
                if _norm_key(island) == _norm_key(dest_name):
                    wh_rates = wh_dist.get(island) or {}
                    costs[cost_key(block, block.columns[0])] = wh_rates.get("minimum")
                    costs[cost_key(block, block.columns[1])] = wh_rates.get("per_kg")
                    break

        # Customs (flat on all rows when value present)
        if isinstance(customs, (int, float)):
            _set_flat_cost(costs, customs_block, float(customs))
        elif customs:
            text_costs[cost_key(customs_block, customs_block.columns[0])] = str(customs)

        # Shipment insurance completas
        if insurance_cc:
            text_costs[cost_key(insurance_cc_block, insurance_cc_block.columns[0])] = str(
                insurance_cc
            )

        # FCL container costs (CARGAS COMPLETAS - CONTENEDORES)
        if origin_zone == ORIGIN_MAINLAND_SPAIN and dest_zone in FCL_ORIGINS:
            _assign_fcl_service_costs(
                costs,
                blocks_by_title,
                coral,
                "Cargas EN CORAL",
                port_name,
                FCL_SIZE_BRACKETS,
            )
            _assign_fcl_service_costs(
                costs,
                blocks_by_title,
                celra,
                "Cargas EN CELRA",
                port_name,
                FCL_SIZE_BRACKETS,
            )
            _assign_fcl_service_costs(
                costs,
                blocks_by_title,
                combo,
                "CARGA EN CELRA / CORAL",
                port_name,
                ("40 pies",),
            )

            if combo.get((port_name, "40 pies")):
                text_costs[
                    cost_key(insurance_fcl_block, insurance_fcl_block.columns[0])
                ] = INSURANCE_FCL_40_LABEL

        # Grupajes from Barcelona
        if origin_zone == ORIGIN_BARCELONA:
            g = None
            for key, values in grupajes.items():
                if _norm_key(key) == _norm_key(dest_name):
                    g = values
                    break
            if g:
                costs[cost_key(grupajes_block, grupajes_block.columns[0])] = g.get("per_kg")
                costs[cost_key(grupajes_block, grupajes_block.columns[1])] = g.get("minimum")

        matrix_rows.append(MatrixRow(shipment=shipment, costs=costs, text_costs=text_costs))

    for section_name, _ in _iter_sections(main):
        if _norm_key(section_name) in {
            _norm_key("RGA EN CORAL"),
            _norm_key("CARGA EN CERLA"),
            _norm_key("Transporte desde barcelona"),
        }:
            skipped.append(
                SkippedItem(
                    source="MainCosts",
                    label=section_name,
                    detail="Section not mapped to matrix columns in this build.",
                )
            )

    return matrix_rows, cost_blocks, skipped


def write_rates_sheet(
    worksheet,
    matrix_rows: list[MatrixRow],
    cost_blocks: list[CostBlock],
) -> None:
    shipment_count = len(SHIPMENT_HEADERS)

    def write_merged_row(row_index: int, values: list[str]) -> None:
        column_index = shipment_count + 1
        for block_index, block in enumerate(cost_blocks):
            width = block_column_width(block)
            value = values[block_index] if block_index < len(values) else ""
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.font = BOLD
            cell.fill = HEADER_FILL
            cell.alignment = LEFT
            if width > 1:
                worksheet.merge_cells(
                    start_row=row_index,
                    start_column=column_index,
                    end_row=row_index,
                    end_column=column_index + width - 1,
                )
            column_index += width

    write_merged_row(COST_NAME_ROW, [block.title for block in cost_blocks])
    write_merged_row(APPLY_IF_ROW, [block.apply_if for block in cost_blocks])
    write_merged_row(RATE_BY_ROW, [block.rate_by for block in cost_blocks])

    for col_idx, header in enumerate(SHIPMENT_HEADERS, start=1):
        cell = worksheet.cell(row=COLUMN_HEADER_ROW, column=col_idx, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = LEFT

    column_index = shipment_count + 1
    for block in cost_blocks:
        if block.uses_shared_currency:
            currency_cell = worksheet.cell(row=BRACKET_ROW, column=column_index)
            currency_cell.fill = HEADER_FILL
            worksheet.cell(row=COLUMN_HEADER_ROW, column=column_index, value=CURRENCY)
            for offset, spec in enumerate(block.columns):
                spec_col = column_index + 1 + offset
                bracket_cell = worksheet.cell(
                    row=BRACKET_ROW, column=spec_col, value=spec.bracket_label
                )
                bracket_cell.font = BOLD
                bracket_cell.fill = HEADER_FILL
                bracket_cell.alignment = CENTER
                unit_cell = worksheet.cell(
                    row=COLUMN_HEADER_ROW, column=spec_col, value=spec.rate_unit
                )
                unit_cell.font = BOLD
                unit_cell.fill = HEADER_FILL
                unit_cell.alignment = CENTER
            column_index += block_column_width(block)
            continue

        for spec in block.columns:
            bracket_cell = worksheet.cell(
                row=BRACKET_ROW, column=column_index, value=spec.bracket_label
            )
            bracket_cell.font = BOLD
            bracket_cell.fill = HEADER_FILL
            bracket_cell.alignment = CENTER
            worksheet.merge_cells(
                start_row=BRACKET_ROW,
                start_column=column_index,
                end_row=BRACKET_ROW,
                end_column=column_index + 1,
            )
            worksheet.cell(row=COLUMN_HEADER_ROW, column=column_index, value=CURRENCY)
            unit_cell = worksheet.cell(
                row=COLUMN_HEADER_ROW, column=column_index + 1, value=spec.rate_unit
            )
            unit_cell.font = BOLD
            unit_cell.fill = HEADER_FILL
            unit_cell.alignment = CENTER
            column_index += 2

    for row_offset, matrix_row in enumerate(matrix_rows):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, header in enumerate(SHIPMENT_HEADERS, start=1):
            worksheet.cell(
                row=excel_row, column=col_idx, value=matrix_row.shipment.get(header)
            )

        column_index = shipment_count + 1
        for block in cost_blocks:
            if block.uses_shared_currency:
                has_data = _block_has_row_data(matrix_row, block)
                if has_data:
                    worksheet.cell(row=excel_row, column=column_index, value=CURRENCY)
                for offset, spec in enumerate(block.columns):
                    spec_col = column_index + 1 + offset
                    key = cost_key(block, spec)
                    if key in matrix_row.text_costs:
                        worksheet.cell(
                            row=excel_row,
                            column=spec_col,
                            value=matrix_row.text_costs[key],
                        )
                        continue
                    value = matrix_row.costs.get(key)
                    if _has_cost(value) and isinstance(value, (int, float)):
                        cell = worksheet.cell(row=excel_row, column=spec_col, value=value)
                        cell.number_format = _excel_number_format(float(value))
                column_index += block_column_width(block)
                continue

            for spec in block.columns:
                key = cost_key(block, spec)
                if key in matrix_row.text_costs:
                    worksheet.cell(
                        row=excel_row,
                        column=column_index,
                        value=matrix_row.text_costs[key],
                    )
                    column_index += 2
                    continue
                value = matrix_row.costs.get(key)
                if _has_cost(value) and isinstance(value, (int, float)):
                    worksheet.cell(row=excel_row, column=column_index, value=CURRENCY)
                    cell = worksheet.cell(
                        row=excel_row, column=column_index + 1, value=value
                    )
                    cell.number_format = _excel_number_format(float(value))
                column_index += 2

    for col_idx in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = 20


def write_skipped_sheet(worksheet, items: list[SkippedItem]) -> None:
    headers = ["Source", "Label", "Detail"]
    for col_idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_idx, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for row_idx, item in enumerate(items, start=2):
        worksheet.cell(row=row_idx, column=1, value=item.source)
        worksheet.cell(row=row_idx, column=2, value=item.label)
        worksheet.cell(row=row_idx, column=3, value=item.detail)
    for col_idx in range(1, 4):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = 36


def build_matrix_workbook(cleaned_path: Path, output_path: Path | None = None) -> Path:
    fields = _load_fields(cleaned_path)
    matrix_rows, cost_blocks, skipped = build_matrix(fields)
    if not matrix_rows:
        raise ValueError("No matrix rows produced.")

    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        stem = cleaned_path.stem.replace(".cleaned", "")
        output_path = OUTPUT_DIR / f"matrix_{stem}_{timestamp}.xlsx"

    workbook = Workbook()
    rates_sheet = workbook.active
    rates_sheet.title = "Rates"
    write_rates_sheet(rates_sheet, matrix_rows, cost_blocks)

    skipped_sheet = workbook.create_sheet("did not added")
    write_skipped_sheet(skipped_sheet, skipped)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Almazar rate matrix XLSX.")
    parser.add_argument(
        "cleaned",
        nargs="?",
        type=Path,
        help="Cleaned JSON path (default: newest *.cleaned.json in processing/)",
    )
    parser.add_argument("-o", "--output", type=Path, help="Output XLSX path")
    args, _unknown = parser.parse_known_args()

    if args.cleaned:
        cleaned_path = args.cleaned
    else:
        candidates = sorted(
            PROCESSING_DIR.glob("*.cleaned.json"),
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
        if not candidates:
            raise FileNotFoundError(f"No cleaned JSON in {PROCESSING_DIR}")
        cleaned_path = candidates[0]

    output_path = build_matrix_workbook(cleaned_path, output_path=args.output)
    print(f"Matrix written to {output_path}")


if __name__ == "__main__":
    main()
